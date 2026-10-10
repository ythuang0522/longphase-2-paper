#!/usr/bin/env python3
"""Issue #5: MethPhaser 0.0.4 against its input and against LongPhase 2 SNV + 5mC co-phasing,
replicate 1, 10-60x. Writes methphaser_vs_longphase.tsv and, with --xlsx, the MethPhaser sheet
of supplementary.xlsx (SNV_Detail columns plus assessed SNV pairs).

MethPhaser and its input are read from compare/ (longphase compare, release v2.1, v5.0q
chr1-22, no BED). LongPhase 2 SNV + 5mC is read from the existing compare TSVs of the paper
runs on the lab server (the same files as supplementary.xlsx co-phase_Detail, case Mod)."""
import os, sys

HERE = os.path.dirname(os.path.abspath(__file__))
R = '/disk/research'
COVS = (10, 20, 30, 40, 50, 60)
RUNS = {
    'MethPhaser input (LongPhase 2 SNV, no GNN)': os.path.join(HERE, 'compare', 'sw_input_{c}x_1.tsv'),
    'MethPhaser 0.0.4': os.path.join(HERE, 'compare', 'sw_methphaser_{c}x_1.tsv'),
    'LongPhase 2 SNV + 5mC, no GNN': R + '/longphase_coh_methyl/sw_longphase_coh_methyl_{c}x_1.tsv',
    'LongPhase 2 SNV + 5mC, GNN': R + '/longphase_coh_methyl_gnn/sw_longphase_coh_methyl_gnn_{c}x_1.tsv',
}


def read(path):
    """Summary line plus the sum of all_assessed_pairs over chromosomes."""
    total, pairs = None, 0
    for line in open(path):
        t = line.rstrip('\n').split('\t')
        if line.startswith('###') and 'Sample' not in line:
            total = t
        elif not line.startswith('#') and len(t) > 11:
            pairs += int(t[11])
    return dict(Phased_SNV=int(total[1]), Phased_SNV_pct=float(total[2]), SNV_SW=int(total[5]),
                SNV_SW_pct=float(total[6]), Hamming_pct=float(total[7]), Blocks=int(total[8]),
                N50=int(total[9]), Block_Sum=int(total[10]), assessed_pairs=pairs)


rows = [(c, run, read(p.format(c=c))) for c in COVS for run, p in RUNS.items()]
with open(os.path.join(HERE, 'methphaser_vs_longphase.tsv'), 'w') as f:
    f.write('coverage\trun\tPhased_SNV\tPhased_SNV(%)\tSNV_SW\tassessed_SNV_pairs\tSNV_SW(%)\t'
            'Hamming(%)\tBlocks\tBlock_N50(bp)\n')
    for c, run, d in rows:
        f.write('\t'.join(map(str, [f'{c}x', run, d['Phased_SNV'], d['Phased_SNV_pct'], d['SNV_SW'],
                                    d['assessed_pairs'], d['SNV_SW_pct'], d['Hamming_pct'],
                                    d['Blocks'], d['N50']])) + '\n')

if '--xlsx' in sys.argv:
    import openpyxl
    path = sys.argv[sys.argv.index('--xlsx') + 1]
    wb = openpyxl.load_workbook(path)
    if 'MethPhaser' in wb.sheetnames:
        del wb['MethPhaser']
    ws = wb.create_sheet('MethPhaser')
    ws.append([None, None, None, 'methphaser_v0.0.4'])
    ws.append(['No.', 'Platform', 'Case', 'Phased_SNV', 'Phased_SNV(%)', 'SNV_SW', 'SNV_SW(%)',
               'HM_Dis.(%)', 'No. Block', 'Block_N50(bp)', 'Block_Sum', 'SNV_assessed_pairs'])
    for i, (c, run, d) in enumerate([r for r in rows if r[1] == 'MethPhaser 0.0.4'], 1):
        ws.append([i, 'ONT R10.4.1', f'{c}x_1', d['Phased_SNV'], d['Phased_SNV_pct'], d['SNV_SW'],
                   d['SNV_SW_pct'], d['Hamming_pct'], d['Blocks'], d['N50'], d['Block_Sum'],
                   d['assessed_pairs']])
    ws.freeze_panes = 'D3'
    wb.save(path)

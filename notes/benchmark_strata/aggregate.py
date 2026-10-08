"""Collect score_phase.py results for issue #1.

Writes
  step0_check.tsv   every run's totals against the manuscript tables
  item4_strata.tsv  SNV switch errors per tool, run, truth, region, stratum
  item4_overlap.tsv switch-error positions shared between tools
  item3_indel.tsv   indel phase accuracy per co-phasing configuration and run
and prints summaries (means over replicates at 10-20x, single run at 30-60x).
"""
import glob
import sys
import os
import re
import statistics as st
from collections import defaultdict

import openpyxl

PAPER = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', '..'))
OUT = sys.argv[1] if len(sys.argv) > 1 else 'out'
COVS = [10, 12, 14, 16, 18, 20, 30, 40, 50, 60]


def read(path):
    rows = {}
    with open(path) as f:
        h = f.readline().lstrip('#').split('\t')
        ham = {h[i]: int(h[i + 1]) for i in range(0, len(h) - 1, 2)}
        cols = f.readline().rstrip('\n').split('\t')
        for line in f:
            c = line.rstrip('\n').split('\t')
            rows[(c[0], c[1])] = dict(zip(cols[2:], map(int, c[2:])))
    return ham, rows


def split_run(name):
    m = re.match(r'(.+)_(\d+)x_(\d+)$', name)
    return m.group(1), int(m.group(2)), int(m.group(3))


# ---------- reference values from the manuscript tables ----------
def table_txt(path):
    ref = {}
    for line in open(path):
        if line.startswith('###') and not line.startswith('###Sample'):
            c = line[3:].rstrip('\n').split('\t')
            name = re.sub(r'(\.phased\.VCF|\.vcf\.gz|\.vcf)$', '', c[0])
            ref[name] = (int(c[1]), int(c[3]), int(c[5]), float(c[7]))
    return ref


def table_xlsx():
    wb = openpyxl.load_workbook(f'{PAPER}/supplementary.xlsx', data_only=True)
    ref = {}
    rows = [r for r in wb['SNV_Detail'].iter_rows(values_only=True) if r[2] and 'x_' in str(r[2])]
    for r in rows:
        if r[1].startswith('ONT'):
            ref[f'longphase_{r[2]}'] = (r[3], 0, r[5], r[7])
    cfg = {'Indel': 'coh_indel', 'Indel+SV': 'coh_indel_sv', 'Mod+Indel': 'coh_indel_methyl',
           'Indel+Mod+SV': 'cophasing'}
    for r in wb['co-phase_Detail'].iter_rows(values_only=True):
        if r[2] in cfg and r[3] and 'x_' in str(r[3]):
            c = cfg[r[2]]
            ref[f'longphase_{c}_{r[3]}'] = (r[4], r[6], r[8], r[10])
            ref[f'longphase_{c}_gnn_{r[3]}'] = (r[14], r[16], r[18], r[20])
            if r[24] is not None:
                ref[f'whatshap_v28_{r[3]}'] = (r[24], r[26], r[28], r[30])
    return ref


REF = {'v5': {**table_xlsx(), **table_txt(f'{PAPER}/20260924_v50q.txt'),
              **table_txt(f'{PAPER}/20260924_cophase_v50q.txt')},
       'v4': table_txt(f'{PAPER}/20260924_GIAB421.txt')}

# ---------- load all results ----------
res = {}   # (item, truth, name) -> (ham, rows)
for p in glob.glob(f'{OUT}/*/*/*.tsv'):
    if p.endswith('.sw.tsv'):
        continue
    item, truth, fn = p.split('/')[-3:]
    res[(item, truth, fn[:-4])] = read(p)

# ---------- step 0 ----------
n_ok = n_bad = n_noref = 0
with open('step0_check.tsv', 'w') as f:
    f.write('item\ttruth\trun\tphased_snv\tphased_indel\tsnv_sw\thamming%\t'
            'table_phased_snv\ttable_phased_indel\ttable_snv_sw\ttable_hamming%\tmatch\n')
    for (item, truth, name), (ham, rows) in sorted(res.items()):
        a = rows[('all', 'any')]
        hm = 100 * ham['ham_all'] / max(1, ham['ham_n'])
        mine = (a['phased_snv'], a['phased_indel'], a['snv_sw'], hm)
        ref = REF[truth].get(name)
        if ref is None:
            ok = 'no_table_value'; n_noref += 1
        else:
            # the SNV-only tables were scored with --only-snvs for some tools,
            # which changes Hamming% and indel counts but not SNV counts
            ok = 'yes' if (mine[0] == ref[0] and mine[2] == ref[2]) else 'NO'
            n_ok += ok == 'yes'; n_bad += ok == 'NO'
        f.write('\t'.join(map(str, [item, truth, name, *mine[:3], f'{hm:.5f}',
                                    *(ref if ref else ('', '', '', '')), ok])) + '\n')
print(f'step 0: {n_ok} runs match the tables (phased SNVs and SNV switch errors), '
      f'{n_bad} do not, {n_noref} have no table value')

# ---------- item 4 ----------
ITEM4 = ['longphase_gnn', 'longphase', 'whatshap_v28_onlySNVs', 'hapcut2_v134']
with open('item4_strata.tsv', 'w') as f:
    f.write('tool\tcoverage\treplicate\ttruth\tregion\tstratum\tphased_snv\tsnv_pairs\tsnv_sw\tsnv_sw_rate%\n')
    for (item, truth, name), (ham, rows) in sorted(res.items()):
        if item != 'item4':
            continue
        tool, cov, rep = split_run(name)
        for (rg, sn), d in rows.items():
            rate = 100 * d['snv_sw'] / d['snv_pairs'] if d['snv_pairs'] else float('nan')
            f.write(f"{tool}\t{cov}\t{rep}\t{truth}\t{rg}\t{sn}\t{d['phased_snv']}\t"
                    f"{d['snv_pairs']}\t{d['snv_sw']}\t{rate:.5f}\n")


def mean_cell(item, tool, truth, cov, rg, sn, key):
    v = [res[(item, truth, f'{tool}_{cov}x_{r}')][1][(rg, sn)][key]
         for r in range(1, 11) if (item, truth, f'{tool}_{cov}x_{r}') in res]
    return st.mean(v) if v else float('nan')


def rate_cell(item, tool, truth, cov, rg, sn, num='snv_sw', den='snv_pairs'):
    """Rate of the replicate sums (pooled), in %."""
    n = d = 0
    for r in range(1, 11):
        k = (item, truth, f'{tool}_{cov}x_{r}')
        if k in res:
            n += res[k][1][(rg, sn)][num]; d += res[k][1][(rg, sn)][den]
    return 100 * n / d if d else float('nan')


short = {'longphase_gnn': 'LP2', 'longphase': 'LP2-noGNN',
         'whatshap_v28_onlySNVs': 'WH', 'hapcut2_v134': 'HC2'}
print('\nItem 4 - SNV switch errors (mean per genome) by benchmark region; fold = max/min of LP2, WH, HC2 rates')
for truth in ('v5', 'v4'):
    for rg in ('all', 'shared', 'v5only', 'v4only', 'neither'):
        print(f'\n[{truth} truth, region {rg}]')
        print('cov\t' + '\t'.join(f'{short[t]}_sw\t{short[t]}_rate%' for t in ITEM4) + '\tfold')
        for cov in COVS:
            cells, rates = [], []
            for t in ITEM4:
                sw = mean_cell('item4', t, truth, cov, rg, 'any', 'snv_sw')
                rt = rate_cell('item4', t, truth, cov, rg, 'any')
                cells += [f'{sw:.0f}', f'{rt:.4f}']
                if t != 'longphase':
                    rates.append(rt)
            fold = max(rates) / min(rates) if min(rates) > 0 else float('nan')
            print(f'{cov}\t' + '\t'.join(cells) + f'\t{fold:.2f}')

print('\nItem 4 - strata inside the v5.0q BED, v5 truth (mean SNV switch errors per genome / pooled rate %)')
for cov in (10, 20, 30, 60):
    print(f'\n[{cov}x]')
    print('stratum\t' + '\t'.join(f'{short[t]}_sw\t{short[t]}_rate%' for t in ITEM4))
    for sn in ['any', 'segdup', 'tandem_repeat', 'homopolymer', 'satellite', 'lowmap_segdup', 'not_difficult']:
        cells = []
        for t in ITEM4:
            cells += [f"{mean_cell('item4', t, 'v5', cov, 'v5bed', sn, 'snv_sw'):.0f}",
                      f"{rate_cell('item4', t, 'v5', cov, 'v5bed', sn):.4f}"]
        print(sn + '\t' + '\t'.join(cells))

# ---------- item 4 overlap of switch positions ----------
def sw_set(truth, tool, run):
    p = f'{OUT}/item4/{truth}/{tool}_{run}.sw.tsv'
    s = set()
    for line in open(p):
        c = line.split('\t')
        if c[3].strip() == 'snv':
            s.add((c[0], int(c[1]), int(c[2])))
    return s


def site_set(truth, tool, run):
    """Switch errors as the set of variant positions flanking them."""
    s = set()
    for ch, a, b in sw_set(truth, tool, run):
        s.add((ch, a)); s.add((ch, b))
    return s


with open('item4_overlap.tsv', 'w') as f:
    f.write('truth\trun\tLP2\tWH\tHC2\tLP2&WH\tLP2&HC2\tWH&HC2\tall3\tWH&HC2_not_LP2\n')
    print('\nItem 4 - switch errors at the same pair of SNVs, replicate 1 (exact pair match)')
    print('truth\trun\tLP2\tWH\tHC2\tLP2&WH\tLP2&HC2\tWH&HC2\tall3\tWH&HC2 only')
    for truth in ('v5', 'v4'):
        for cov in COVS:
            run = f'{cov}x_1'
            L, W, H = (sw_set(truth, t, run) for t in ('longphase_gnn', 'whatshap_v28_onlySNVs', 'hapcut2_v134'))
            row = [truth, run, len(L), len(W), len(H), len(L & W), len(L & H), len(W & H),
                   len(L & W & H), len((W & H) - L)]
            f.write('\t'.join(map(str, row)) + '\n')
            print('\t'.join(map(str, row)))

# ---------- item 3 ----------
CFG = ['longphase_coh_indel', 'longphase_coh_indel_gnn', 'longphase_coh_indel_sv',
       'longphase_coh_indel_sv_gnn', 'longphase_coh_indel_methyl', 'longphase_coh_indel_methyl_gnn',
       'longphase_cophasing', 'longphase_cophasing_gnn', 'whatshap_v28']
with open('item3_indel.tsv', 'w') as f:
    f.write('config\tcoverage\treplicate\tregion\tphased_snv\tphased_indel\tsnv_pairs\tsnv_sw\t'
            'all_pairs\tall_sw\tindel_pairs\tindel_sw\tindel_sw_rate%\tsnv_indel_sw_rate%\t'
            'indel_hamming\tindels_in_blocks\tindel_hamming%\tsnv_sw_rate%\n')
    for (item, truth, name), (ham, rows) in sorted(res.items()):
        if item != 'item3':
            continue
        cfg, cov, rep = split_run(name)
        for rg in ('all', 'v5bed'):
            d = rows[(rg, 'any')]
            ih = (f"{ham['ham_indel']}\t{ham['ham_n_indel']}\t"
                  f"{100 * ham['ham_indel'] / max(1, ham['ham_n_indel']):.4f}") if rg == 'all' else '\t\t'
            f.write(f"{cfg}\t{cov}\t{rep}\t{rg}\t{d['phased_snv']}\t{d['phased_indel']}\t{d['snv_pairs']}\t"
                    f"{d['snv_sw']}\t{d['all_pairs']}\t{d['all_sw']}\t{d['indel_pairs']}\t{d['indel_sw']}\t"
                    f"{100 * d['indel_sw'] / max(1, d['indel_pairs']):.4f}\t"
                    f"{100 * d['all_sw'] / max(1, d['all_pairs']):.4f}\t{ih}\t"
                    f"{100 * d['snv_sw'] / max(1, d['snv_pairs']):.4f}\n")

print('\nItem 3 - indel phase accuracy against v5.0q (no BED, as in the tables); means per genome')
print('config\tcov\tphased_indel\tindel_pairs\tindel_sw\tindel_sw%\tsnv+indel_sw%\tsnv_sw\tsnv_sw%\tindel_ham%')
for cfg in CFG:
    for cov in COVS:
        def m(key, rg='all'):
            return mean_cell('item3', cfg, 'v5', cov, rg, 'any', key)
        hams = [res[('item3', 'v5', f'{cfg}_{cov}x_{r}')][0] for r in range(1, 11)
                if ('item3', 'v5', f'{cfg}_{cov}x_{r}') in res]
        if not hams:
            continue
        ih = 100 * sum(h['ham_indel'] for h in hams) / max(1, sum(h['ham_n_indel'] for h in hams))
        print(f"{cfg}\t{cov}\t{m('phased_indel'):.0f}\t{m('indel_pairs'):.0f}\t{m('indel_sw'):.0f}\t"
              f"{rate_cell('item3', cfg, 'v5', cov, 'all', 'any', 'indel_sw', 'indel_pairs'):.4f}\t"
              f"{rate_cell('item3', cfg, 'v5', cov, 'all', 'any', 'all_sw', 'all_pairs'):.4f}\t"
              f"{m('snv_sw'):.0f}\t{rate_cell('item3', cfg, 'v5', cov, 'all', 'any'):.4f}\t{ih:.3f}")

#!/usr/bin/env python3
"""Variant_Calling sheet of supplementary.xlsx: Hetero. SV, Homo. SV (Sniffles2 2.8.0, now called with
--reference) and Hetero. methyl (longphase modcall v2.1.2) of the 64 nanopore runs, from the v2.1.2 inputs.
Hetero. SV: GT 0/1, any FILTER; Homo. SV: GT 1/1; Hetero. methyl: GT 0/1 (all contigs). The PEPPER columns
and the HiFi rows (no SV or 5mC columns) are unchanged.

    python update_variant_calling.py check   # the definitions reproduce the current values from the old inputs
    python update_variant_calling.py write
"""
import os, sys
import openpyxl
ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', '..')
X = os.path.join(ROOT, 'supplementary.xlsx')
R = '/disk/research'
SRC = {'check': (R + '/sniffles_v2.8/sniffles_{}.vcf', R + '/methyl/modcall_{}.vcf'),
       'write': (R + '/v2.1.2/sniffles/sniffles_{}.vcf', R + '/v2.1.2/methyl/modcall_{}.vcf')}


def gts(p):
    n = {}
    for l in open(p):
        if l[0] == '#':
            continue
        g = l.split('\t', 10)[9].split(':', 1)[0].strip()
        n[g] = n.get(g, 0) + 1
    return n


mode = sys.argv[1] if len(sys.argv) > 1 else 'check'
wb = openpyxl.load_workbook(X)
ws = wb['Variant_Calling']
head = [c.value for c in ws[2]]
cs, ch, cm = head.index('Hetero. SV') + 1, head.index('Homo. SV') + 1, head.index('Hetero. methyl') + 1
bad = n = 0
for r in range(3, ws.max_row + 1):
    if not str(ws.cell(r, 2).value or '').startswith('ONT'):
        continue
    case = ws.cell(r, 3).value
    sv, mod = (p.format(case) for p in SRC[mode])
    a, b = gts(sv), gts(mod)
    vals = [a.get('0/1', 0), a.get('1/1', 0), b.get('0/1', 0)]
    cur = [ws.cell(r, c).value for c in (cs, ch, cm)]
    if mode == 'check':
        if cur != vals:
            bad += 1; print('differs', case, cur, vals)
    else:
        for c, v in zip((cs, ch, cm), vals):
            ws.cell(r, c).value = v
        n += 1
if mode == 'write':
    ws.cell(1, cs).value = 'Sniffles_v2.8 (--reference)'
    ws.cell(1, cm).value = 'LongPhase modcall v2.1.2'
    wb.save(X); print('rows written:', n)
else:
    print('mismatches:', bad)

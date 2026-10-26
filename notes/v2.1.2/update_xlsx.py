#!/usr/bin/env python3
"""Update supplementary.xlsx (Supplementary Data 1) to LongPhase v2.1.2.

v2.1.2 changes the output only of runs with an SV file or 5mC (SV block rule, new Sniffles calls with
--reference, modcall v2.1.2); SNV-only and SNV + indel output is identical to v2.1.1.

  SNV_Detail       block `longphase_v2.1.1` renamed `longphase_v2.1.2`; values unchanged (identical VCFs)
  co-phase_Detail  block `longphase_v2.1.1` renamed `longphase_v2.1.2`; Mod, SV, Indel+SV, Mod+Indel and
                   Indel+Mod+SV rows of both LongPhase blocks (with GNN correction, and without it in the
                   block still labelled `longphase_v2.0.1`) from the `longphase compare` reports of
                   /disk/research/v2.1.2; the Indel rows are unchanged
  Indel_accuracy   rewritten from notes/benchmark_strata/item3_indel.tsv
  SV_5mC_phasing   notes/sv_mod/build_sheet.py (run separately)

    python update_xlsx.py check    # old values against the v2.1.1 / v2.1 reports (mapping check), no write
    python update_xlsx.py write
"""
import csv, os, sys
import openpyxl

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', '..')
X = os.environ.get('XLSX', os.path.join(ROOT, 'supplementary.xlsx'))
R, OLDG, NEW = '/disk/research', '/disk/research/v2.1.1', '/disk/research/v2.1.2'
CHANGED = {'Mod': 'coh_methyl', 'SV': 'coh_sv', 'Indel+SV': 'coh_indel_sv', 'Mod+Indel': 'coh_indel_methyl',
           'Indel+Mod+SV': 'cophasing'}
TSV_SHEETS = {'Indel_accuracy': 'benchmark_strata/item3_indel.tsv'}


def report(path):
    for l in open(path):
        if l.startswith('###') and not l.startswith('###Sample'):
            t = l.rstrip('\n').split('\t')
            return [int(t[1]), float(t[2]), int(t[3]), float(t[4]), int(t[5]), float(t[6]),
                    float(t[7]), int(t[8]), int(t[9]), int(t[10])]
    raise ValueError(path)


def coph(cfg, case, gnn, new):
    d = 'longphase_' + CHANGED[cfg] + ('_gnn' if gnn else '')
    if new:
        return report(f'{NEW}/{d}/sw_{d}_{case}.tsv')
    if gnn:
        return report(f'{OLDG}/{d}/sw_{d}_{case}.tsv')
    for p in (f'{R}/{d}/sw_{d}_{case}.tsv', f'{R}/{d}/{d}_{case}.tsv'):
        if os.path.exists(p):
            return report(p)
    raise OSError(d + case)


def col(ws, names):
    head = [c.value for c in ws[1]]
    for n in names:
        if n in head:
            return head.index(n) + 1
    raise KeyError(names)


def num(v):
    for f in (int, float):
        try:
            return f(v)
        except ValueError:
            pass
    return None if v in ('', 'NA') else v


def main(mode):
    wb = openpyxl.load_workbook(X)
    bad = changed = 0
    ws = wb['SNV_Detail']
    g = col(ws, ('longphase_v2.1.1', 'longphase_v2.1.2'))
    if mode == 'write':
        ws.cell(1, g).value = 'longphase_v2.1.2'
    ws = wb['co-phase_Detail']
    g = col(ws, ('longphase_v2.1.1', 'longphase_v2.1.2')); n = col(ws, ('longphase_v2.0.1',))
    for r in range(3, ws.max_row + 1):
        cfg, case = ws.cell(r, 3).value, ws.cell(r, 4).value
        if cfg not in CHANGED or not case:
            continue
        for c0, gnn in ((g, True), (n, False)):
            cur = [ws.cell(r, c0 + i).value for i in range(10)]
            if mode == 'check':
                old = coph(cfg, case, gnn, False)
                if [round(float(a), 4) for a in cur] != [round(float(b), 4) for b in old]:
                    bad += 1; print('differs', cfg, case, gnn, cur, old)
            else:
                vals = coph(cfg, case, gnn, True)
                changed += cur != vals
                for i, v in enumerate(vals):
                    ws.cell(r, c0 + i).value = v
    if mode == 'write':
        ws.cell(1, g).value = 'longphase_v2.1.2'
    for sheet, src in TSV_SHEETS.items():
        ws = wb[sheet]
        rows = list(csv.reader(open(os.path.join(ROOT, 'notes', src)), delimiter='\t'))
        if mode == 'check':
            cur = [list(c) for c in ws.iter_rows(values_only=True)]
            same = len(cur) == len(rows) and all([num(x) for x in t] == c[:len(t)] for c, t in zip(cur, rows))
            print(f'{sheet}: {"same as" if same else "DIFFERS from"} {src}'); bad += not same
        else:
            font = ws.cell(1, 1).font.copy()
            ws.delete_rows(1, ws.max_row)
            for i, t in enumerate(rows):
                ws.append([x if i == 0 else num(x) for x in t])
            for row in ws.iter_rows():
                for c in row:
                    c.font = font
    if mode == 'write':
        wb.save(X); print('written', X, '; co-phase_Detail run blocks changed:', changed)
    else:
        print('mismatches:', bad)


if __name__ == '__main__':
    main(sys.argv[1] if len(sys.argv) > 1 else 'check')

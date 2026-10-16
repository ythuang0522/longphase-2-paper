#!/usr/bin/env python3
"""Update supplementary.xlsx (Supplementary Data 1) from LongPhase v2.1.1 (GNN correction) runs.

Only values of LongPhase 2 with GNN correction change; every other tool, and LongPhase 2 without
correction (`longphase_v2.0.1`), is left as it is.

  SNV_Detail, co-phase_Detail   block `longphase_v2.1` -> `longphase_v2.1.1`, values from the
                                `longphase compare` reports of /disk/research/v2.1.1 (nanopore)
                                and /disk/research/v2.1.1/pacbio (HiFi)
  SV_5mC_phasing                notes/sv_mod/build_sheet.py (run separately)
  other sheets                  rewritten from their source TSVs in notes/ (one sheet = one TSV)

    python3 update_xlsx.py check    # old sheets against the v2.1 sources (mapping check), no write
    python3 update_xlsx.py write    # write the v2.1.1 values
"""
import csv, os, sys
import openpyxl

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', '..')
X = os.environ.get('XLSX', os.path.join(ROOT, 'supplementary.xlsx'))
NEW = '/disk/research/v2.1.1'
OLD = '/disk/research'
COPH = {'Indel': 'coh_indel', 'Mod': 'coh_methyl', 'SV': 'coh_sv', 'Indel+SV': 'coh_indel_sv',
        'Mod+Indel': 'coh_indel_methyl', 'Indel+Mod+SV': 'cophasing'}
TSV_SHEETS = {  # sheet -> source TSV (relative to notes/)
    'SNV_switch_by_region': 'benchmark_strata/item4_strata.tsv',
    'SNV_switch_coincidence': 'benchmark_strata/item4_overlap.tsv',
    'Indel_accuracy': 'benchmark_strata/item3_indel.tsv',
    'Unphased_composition': 'unphased/issue6_composition.tsv',
    'Unphased_intervals': 'unphased/issue6_intervals.tsv',
    'Unphased_depth': 'unphased/issue6_depth_hist.tsv',
    'Unphased_sets': 'unphased/unphased_sets.tsv',
    'Phased_calls': 'unphased/phased_calls_chr1_22.tsv',
}


def report(path):
    """### summary line of a longphase compare report -> list of numbers (fields 1..10)."""
    for l in open(path):
        if l.startswith('###') and not l.startswith('###Sample'):
            t = l.rstrip('\n').split('\t')
            return [int(t[1]), float(t[2]), int(t[3]), float(t[4]), int(t[5]), float(t[6]),
                    float(t[7]), int(t[8]), int(t[9]), int(t[10])]
    raise ValueError(path)


def snv_report(base, platform, case):
    if platform.startswith('PB'):
        return report(f'{base}/pacbio/sw_longphase_hifi_{case}.tsv')
    return report(f'{base}/longphase_gnn/sw_longphase_gnn_{case}.tsv')


def coph_report(base, cfg, case):
    d = 'longphase_' + COPH[cfg] + '_gnn'
    return report(f'{base}/{d}/sw_{d}_{case}.tsv')


def block_col(ws, name):
    head = [c.value for c in ws[1]]
    if name not in head and name + '.1' in head:   # already renamed by an earlier write
        name += '.1'
    return head.index(name) + 1  # 1-based column of the block's first value


def num(v):
    try:
        return int(v)
    except ValueError:
        try:
            return float(v)
        except ValueError:
            return None if v in ('', 'NA') else v


def main(mode):
    wb = openpyxl.load_workbook(X)
    bad = 0
    # SNV_Detail: Phased_SNV, %, SNV_SW, SNV_SW(%), HM, No. Block, N50, Block_Sum
    ws = wb['SNV_Detail']
    c0 = block_col(ws, 'longphase_v2.1')
    for r in range(3, ws.max_row + 1):
        plat, case = ws.cell(r, 2).value, ws.cell(r, 3).value
        if not case or 'x_' not in str(case):
            continue
        base = OLD if mode == 'check' else NEW
        try:
            rep = snv_report(base, plat, case)
        except (OSError, ValueError):
            print('MISSING report, row kept:', plat, case); continue
        vals = [rep[0], rep[1], rep[4], rep[5], rep[6], rep[7], rep[8], rep[9]]
        if mode == 'check':
            if plat.startswith('PB'):
                continue  # v2.1 HiFi reports were overwritten in place by the 2026-10 issue #4 rerun
            old = [ws.cell(r, c0 + i).value for i in range(8)]
            if [round(float(a), 4) for a in old] != [round(float(b), 4) for b in vals]:
                bad += 1; print('SNV_Detail differs', plat, case, old, vals)
        else:
            for i, v in enumerate(vals):
                ws.cell(r, c0 + i).value = v
    if mode == 'write':
        ws.cell(1, c0).value = 'longphase_v2.1.1'
    # co-phase_Detail: Phased_SNV, %, Phased_INDEL, %, SNV_SW, SNV_SW(%), HM, No. Block, N50, Block_Sum
    ws = wb['co-phase_Detail']
    c0 = block_col(ws, 'longphase_v2.1')
    for r in range(3, ws.max_row + 1):
        cfg, case = ws.cell(r, 3).value, ws.cell(r, 4).value
        if cfg not in COPH or not case:
            continue
        vals = coph_report(OLD if mode == 'check' else NEW, cfg, case)
        if mode == 'check':
            old = [ws.cell(r, c0 + i).value for i in range(10)]
            if [round(float(a), 4) for a in old] != [round(float(b), 4) for b in vals]:
                bad += 1; print('co-phase_Detail differs', cfg, case, old, vals)
        else:
            for i, v in enumerate(vals):
                ws.cell(r, c0 + i).value = v
    if mode == 'write':
        ws.cell(1, c0).value = 'longphase_v2.1.1'
    # sheets that are one TSV each
    for sheet, src in TSV_SHEETS.items():
        ws = wb[sheet]
        rows = list(csv.reader(open(os.path.join(ROOT, 'notes', src)), delimiter='\t'))
        if mode == 'check':
            cur = [[c for c in row] for row in ws.iter_rows(values_only=True)]
            same = len(cur) == len(rows) and all(
                [str(a) if a is not None else '' for a in c][:len(t)] == t or
                [num(x) for x in t] == [a for a in c][:len(t)] for c, t in zip(cur, rows))
            print(f'{sheet}: {"same as" if same else "DIFFERS from"} {src}')
            bad += not same
        else:
            font = ws.cell(1, 1).font.copy()
            ws.delete_rows(1, ws.max_row)
            for i, t in enumerate(rows):
                ws.append([x if i == 0 else num(x) for x in t])
            for row in ws.iter_rows():
                for c in row:
                    c.font = font
    if mode == 'write':
        wb.save(X)
        print('written', X)
    else:
        print('mismatches:', bad)


if __name__ == '__main__':
    main(sys.argv[1] if len(sys.argv) > 1 else 'check')

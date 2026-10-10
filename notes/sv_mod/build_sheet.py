#!/usr/bin/env python3
"""Write sv_mod.tsv (all contigs, as in Variant_Calling) into the SV_5mC_phasing sheet of
supplementary.xlsx, one row per co-phasing configuration and run, in the column layout of
co-phase_Detail (longphase_v2.0.1 = without GNN correction, longphase_v2.1 = with it)."""
import openpyxl
X = '../../supplementary.xlsx'
rows = [l.rstrip('\n').split('\t') for l in open('sv_mod.tsv')]
h, rows = rows[0], rows[1:]
ix = {k: i for i, k in enumerate(h)}
wb = openpyxl.load_workbook(X)
if 'SV_5mC_phasing' in wb.sheetnames: del wb['SV_5mC_phasing']
ws = wb.create_sheet('SV_5mC_phasing')
ws.append([None] * 6 + ['longphase_v2.0.1', None, None, None, 'longphase_v2.1'])
ws.append(['No.', 'Platform', 'Co-phase', 'Case', 'Hetero. SV', 'Hetero. methyl'] +
          ['Phased_SV', 'Phased_SV(%)', 'Phased_methyl', 'Phased_methyl(%)'] * 2)
num = lambda v: None if v == 'NA' else int(v)
pct = lambda a, b: None if a is None else round(100 * a / b, 5)
for n, t in enumerate(rows, 1):
    g = lambda k: num(t[ix[k]])
    hs, hm = g('het_SV_all'), g('het_5mC_all')
    out = [n, 'ONT R10.4.1', t[0], f'{t[2]}x_{t[3]}', hs, hm]
    for p in ('', 'gnn_'):
        s, m = g(p + 'phased_SV_all'), g(p + 'phased_5mC_all')
        out += [s, pct(s, hs), m, pct(m, hm)]
    ws.append(out)
for row in ws.iter_rows():
    for c in row: c.font = openpyxl.styles.Font(name='Arial', size=10)
ws.freeze_panes = 'E3'
wb.save(X)
print(ws.max_row - 2, 'rows')
